"""Leave reports with JSON / CSV / Excel output."""

import csv
import io
from collections import defaultdict
from datetime import date

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import extract, select
from sqlalchemy.orm import Session, selectinload

from ..auth import get_current_user, is_manager, subordinate_ids
from ..database import get_db
from ..models import Department, Employee, LeaveRequest, LeaveType, RequestStatus
from ..services import balances, workdays

router = APIRouter(prefix="/reports", tags=["reports"])
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def _scope(db: Session, user: Employee) -> set[int] | None:
    """None = unrestricted (HR). Managers are limited to their reporting tree."""
    if user.is_hr:
        return None
    if not is_manager(db, user):
        raise HTTPException(403, "Reports are available to managers and HR")
    return subordinate_ids(db, user.id) | {user.id}


def _employees(db: Session, scope: set[int] | None, department_id: int | None, include_inactive: bool = False) -> list[Employee]:
    stmt = select(Employee).options(selectinload(Employee.department), selectinload(Employee.designation))
    if not include_inactive:
        stmt = stmt.where(Employee.status == "ACTIVE")
    if scope is not None:
        stmt = stmt.where(Employee.id.in_(scope))
    if department_id:
        stmt = stmt.where(Employee.department_id == department_id)
    return list(db.scalars(stmt.order_by(Employee.first_name, Employee.last_name)))


def export(rows: list[dict], columns: list[tuple[str, str]], filename: str, fmt: str, title: str):
    if fmt == "json":
        return {"columns": [{"key": k, "label": label} for k, label in columns], "rows": rows}
    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([label for _, label in columns])
        for r in rows:
            writer.writerow([r.get(k, "") if r.get(k) is not None else "" for k, _ in columns])
        data = io.BytesIO(buf.getvalue().encode("utf-8-sig"))
        return StreamingResponse(
            data, media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="{filename}.csv"'}
        )
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = title[:31]
        ws.append([label for _, label in columns])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4F46E5")
        for r in rows:
            ws.append([r.get(k) for k, _ in columns])
        for i, (k, label) in enumerate(columns, start=1):
            width = max([len(str(label))] + [len(str(r.get(k) or "")) for r in rows[:500]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 50)
        ws.freeze_panes = "A2"
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return StreamingResponse(
            out,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}.xlsx"'},
        )
    raise HTTPException(400, "format must be json, csv or xlsx")


@router.get("/utilization")
def utilization(
    year: int | None = None,
    department_id: int | None = None,
    format: str = "json",
    db: Session = Depends(get_db),
    user: Employee = Depends(get_current_user),
):
    """Per employee, per leave type: credited, used, pending, available and utilization %."""
    year = year or date.today().year
    scope = _scope(db, user)
    rows = []
    for emp in _employees(db, scope, department_id):
        for b in balances.employee_balances(db, emp, year):
            credited = b["credited"]
            rows.append(
                {
                    "employee_code": emp.employee_code,
                    "employee": emp.full_name,
                    "country": emp.country_code or "",
                    "department": emp.department.name if emp.department else "",
                    "leave_type": b["leave_type_name"],
                    "leave_type_code": b["leave_type_code"],
                    "carried_forward": b["carried_forward"],
                    "credited": credited,
                    "used": b["used"],
                    "pending": b["pending"],
                    "available": b["available"],
                    "utilization_pct": round(b["used"] / credited * 100, 1) if credited else None,
                }
            )
    db.commit()
    cols = [
        ("employee_code", "Code"),
        ("employee", "Employee"),
        ("country", "Country"),
        ("department", "Department"),
        ("leave_type", "Leave type"),
        ("leave_type_code", "Type code"),
        ("carried_forward", "Carried forward"),
        ("credited", "Credited"),
        ("used", "Used"),
        ("pending", "Pending"),
        ("available", "Available"),
        ("utilization_pct", "Utilization %"),
    ]
    return export(rows, cols, f"leave-utilization-{year}", format, "Utilization")


@router.get("/employee-history")
def employee_history(
    employee_id: int | None = None,
    year: int | None = None,
    department_id: int | None = None,
    status: str | None = None,
    format: str = "json",
    db: Session = Depends(get_db),
    user: Employee = Depends(get_current_user),
):
    scope = _scope(db, user)
    if employee_id and scope is not None and employee_id not in scope:
        raise HTTPException(403, "Not allowed")
    stmt = (
        select(LeaveRequest)
        .options(
            selectinload(LeaveRequest.employee).selectinload(Employee.department),
            selectinload(LeaveRequest.leave_type),
            selectinload(LeaveRequest.decided_by),
        )
        .join(Employee, Employee.id == LeaveRequest.employee_id)
    )
    if scope is not None:
        stmt = stmt.where(LeaveRequest.employee_id.in_(scope))
    if employee_id:
        stmt = stmt.where(LeaveRequest.employee_id == employee_id)
    if department_id:
        stmt = stmt.where(Employee.department_id == department_id)
    if year:
        stmt = stmt.where(extract("year", LeaveRequest.start_date) == year)
    if status:
        stmt = stmt.where(LeaveRequest.status == status)
    rows = [
        {
            "id": r.id,
            "employee_code": r.employee.employee_code,
            "employee": r.employee.full_name,
            "department": r.employee.department.name if r.employee.department else "",
            "leave_type": r.leave_type.name,
            "start_date": r.start_date.isoformat(),
            "end_date": r.end_date.isoformat(),
            "duration_type": r.duration_type.replace("_", " ").title(),
            "total_days": float(r.total_days),
            "status": r.status,
            "reason": r.reason or "",
            "applied_at": r.applied_at.strftime("%Y-%m-%d %H:%M"),
            "decided_by": r.decided_by.full_name if r.decided_by else "",
            "decision_comment": r.decision_comment or "",
        }
        for r in db.scalars(stmt.order_by(LeaveRequest.start_date.desc()))
    ]
    cols = [
        ("employee_code", "Code"),
        ("employee", "Employee"),
        ("department", "Department"),
        ("leave_type", "Leave type"),
        ("start_date", "From"),
        ("end_date", "To"),
        ("duration_type", "Duration"),
        ("total_days", "Days"),
        ("status", "Status"),
        ("reason", "Reason"),
        ("applied_at", "Applied at"),
        ("decided_by", "Decided by"),
        ("decision_comment", "Comment"),
    ]
    return export(rows, cols, f"leave-history-{year or 'all'}", format, "Leave history")


def _approved(db: Session, year: int, scope: set[int] | None) -> list[LeaveRequest]:
    stmt = (
        select(LeaveRequest)
        .options(selectinload(LeaveRequest.employee).selectinload(Employee.department), selectinload(LeaveRequest.leave_type))
        .where(LeaveRequest.status == RequestStatus.APPROVED.value, extract("year", LeaveRequest.start_date) == year)
    )
    if scope is not None:
        stmt = stmt.where(LeaveRequest.employee_id.in_(scope))
    return list(db.scalars(stmt))


def _days_by_month(db: Session, r: LeaveRequest) -> dict[int, float]:
    """Split a leave's cost across the months its working days fall in."""
    if r.duration_type != "FULL_DAY" or r.start_date.month == r.end_date.month:
        return {r.start_date.month: float(r.total_days)}
    out: dict[int, float] = defaultdict(float)
    for d in workdays.working_dates(db, r.start_date, r.end_date, r.employee_id):
        out[d.month] += 1
    return out


@router.get("/department")
def department_report(
    year: int | None = None,
    format: str = "json",
    db: Session = Depends(get_db),
    user: Employee = Depends(get_current_user),
):
    year = year or date.today().year
    scope = _scope(db, user)
    types = list(db.scalars(select(LeaveType).where(LeaveType.is_active.is_(True)).order_by(LeaveType.sort_order)))
    depts = {d.id: d.name for d in db.scalars(select(Department))}
    headcount: dict[str, int] = defaultdict(int)
    for e in _employees(db, scope, None):
        headcount[depts.get(e.department_id, "Unassigned")] += 1
    days: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    for r in _approved(db, year, scope):
        dname = r.employee.department.name if r.employee.department else "Unassigned"
        days[dname][r.leave_type.code] += float(r.total_days)
    rows = []
    for dname in sorted(set(headcount) | set(days)):
        total = sum(days[dname].values())
        row = {"department": dname, "headcount": headcount.get(dname, 0), "total_days": round(total, 2)}
        row["avg_per_employee"] = round(total / row["headcount"], 2) if row["headcount"] else None
        for lt in types:
            row[lt.code] = round(days[dname].get(lt.code, 0), 2)
        rows.append(row)
    cols = [("department", "Department"), ("headcount", "Headcount")]
    cols += [(lt.code, lt.name) for lt in types]
    cols += [("total_days", "Total days"), ("avg_per_employee", "Avg / employee")]
    return export(rows, cols, f"department-leave-{year}", format, "Departments")


@router.get("/monthly")
def monthly_report(
    year: int | None = None,
    month: int | None = None,
    format: str = "json",
    db: Session = Depends(get_db),
    user: Employee = Depends(get_current_user),
):
    """Without `month`: leave days per month by type for the year. With `month`: per-employee breakdown."""
    year = year or date.today().year
    scope = _scope(db, user)
    types = list(db.scalars(select(LeaveType).where(LeaveType.is_active.is_(True)).order_by(LeaveType.sort_order)))
    leaves = _approved(db, year, scope)

    if month is None:
        grid: dict[int, dict[str, float]] = defaultdict(lambda: defaultdict(float))
        for r in leaves:
            for m, d in _days_by_month(db, r).items():
                grid[m][r.leave_type.code] += d
        rows = []
        for m in range(1, 13):
            row = {"month": f"{MONTHS[m - 1]} {year}", "month_num": m}
            for lt in types:
                row[lt.code] = round(grid[m].get(lt.code, 0), 2)
            row["total_days"] = round(sum(grid[m].values()), 2)
            rows.append(row)
        cols = [("month", "Month")] + [(lt.code, lt.name) for lt in types] + [("total_days", "Total days")]
        return export(rows, cols, f"monthly-leave-{year}", format, "Monthly")

    if not 1 <= month <= 12:
        raise HTTPException(400, "month must be 1-12")
    per_emp: dict[int, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    for r in leaves:
        d = _days_by_month(db, r).get(month)
        if d:
            per_emp[r.employee_id][r.leave_type.code] += d
    rows = []
    for emp in _employees(db, scope, None, include_inactive=True):
        if emp.id not in per_emp and emp.status != "ACTIVE":
            continue
        row = {
            "employee_code": emp.employee_code,
            "employee": emp.full_name,
            "department": emp.department.name if emp.department else "",
        }
        for lt in types:
            row[lt.code] = round(per_emp[emp.id].get(lt.code, 0), 2)
        row["total_days"] = round(sum(per_emp[emp.id].values()), 2)
        rows.append(row)
    cols = [("employee_code", "Code"), ("employee", "Employee"), ("department", "Department")]
    cols += [(lt.code, lt.name) for lt in types] + [("total_days", "Total days")]
    return export(rows, cols, f"leave-{year}-{month:02d}", format, f"{MONTHS[month - 1]} {year}")
